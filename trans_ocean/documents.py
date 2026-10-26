"""Document specialist agent (stage 3).

- Keeps company requisites ("Реквизиты компании") and company documents with expiry dates
  ("Документы"); the Coordinator's morning digest already warns about documents expiring
  within 30 days.
- /documents in the management group: status of company documents and what is missing for tenders.
- /paket <лот> [цена]: builds a draft tender package (Word files, in Uzbek) from the tender row
  in "Тендеры" and the requisites, sends it to the group and creates a task. A person checks,
  signs and uploads; nothing is submitted by the bot.
"""

from __future__ import annotations

import html
import io
import re
from datetime import datetime, timedelta
from typing import Any, Callable, Dict, List, Optional, Tuple

from .base import Base, now_local, parse_date, stamp

AGENT = "Документовед"
REQ_SHEET = "Реквизиты компании"
DOC_SHEET = "Документы"
TENDER_SHEETS = ("Тендеры", "Тендеры (архив)")

DEFAULT_REQUISITES = [
    ("Название", "«TRANS OCEAN LOGISTICS» МЧЖ", "Как в свидетельстве"),
    ("Директор", "Нуриддинов Тўлаган Акоидинович", "Полное ФИО"),
    ("Директор кратко", "Нуриддинов Т. А.", ""),
    ("СТИР", "303372643", "ИНН"),
    ("ОКЭД", "49410", ""),
    ("Адрес", "Тошкент ш., Яккасарой тумани, Шота Руставели кўчаси, 45-уй", "Юридический адрес"),
    ("Банк", "«Ўзмиллийбанк» АЖ Учтепа филиали", "Проверить у бухгалтера"),
    ("МФО", "00450", "Проверить у бухгалтера"),
    ("Расчётный счёт", "20208000600477625001", "Проверить у бухгалтера"),
    ("Телефон", "+998 98 305-72-34, +998 99 898-12-34", ""),
    ("E-mail", "ocean_logistics@mail.ru", ""),
    ("Свидетельство", "№ 007282, 22.06.2015 й.", "Давлат рўйхатидан ўтказилганлиги тўғрисидаги гувоҳнома"),
    ("Лицензия", "АТ № 2000764, 11.02.2020 й. (муддатсиз)", "Вид: городские и пригородные перевозки грузов"),
    ("Парк", "20 бирлик: паст рамали тиркамалар (трал) 1,5–50 т, тентли ярим тиркамалар 96 м³, контейнер ташувчилар, бортли ярим тиркамалар", ""),
    ("Техосмотр и медосмотр", "«AZIZ-BAXT» МЧЖ, шартнома № 28, 01.01.2026 й.", "Предрейсовый техосмотр машин и медосмотр водителей"),
    ("Стоянка", "Тошкент вилояти, Тошкент тумани, Илғор кўчаси, 13а-уй, 2100 м²", "Аренда до 16.12.2026"),
    ("Подача машины по Ташкенту (часов)", "4", "Обязательство в техническом предложении"),
    ("Подача машины по регионам (часов)", "24", "Обязательство в техническом предложении"),
    ("Начало работ после договора (рабочих дней)", "1", ""),
    ("Опыт водителей (лет)", "3", ""),
    ("Страховая компания", "", "ОСГО и страхование грузов — вписать"),
    ("Партнёры (самосвалы, фургоны, манипуляторы)", "", "Названия партнёров-перевозчиков"),
]

DEFAULT_DOCUMENTS = [
    {"Документ": "Свидетельство о госрегистрации", "Тип": "Учредительный", "Номер": "№ 007282",
     "Дата выдачи": "2015-06-22", "Действует до": "", "Заметки": "бессрочно"},
    {"Документ": "Устав (новая редакция)", "Тип": "Учредительный", "Номер": "",
     "Дата выдачи": "2021", "Действует до": "", "Заметки": "бессрочно"},
    {"Документ": "Лицензия на перевозку грузов", "Тип": "Лицензия", "Номер": "АТ № 2000764",
     "Дата выдачи": "2020-02-11", "Действует до": "", "Заметки": "бессрочно; городские и пригородные перевозки"},
    {"Документ": "Аренда офиса (Шота Руставели 45)", "Тип": "Договор", "Номер": "№ 10/2026, ID 3546857",
     "Дата выдачи": "2026-01-01", "Действует до": "2026-12-31", "Заметки": "Исхаков А.К."},
    {"Документ": "Техосмотр машин и медосмотр водителей", "Тип": "Договор", "Номер": "№ 28 (AZIZ-BAXT)",
     "Дата выдачи": "2026-01-01", "Действует до": "2026-12-31", "Заметки": ""},
    {"Документ": "Аренда стоянки 2100 м² (KELES RENESANS)", "Тип": "Договор", "Номер": "ID 4244404",
     "Дата выдачи": "2026-07-17", "Действует до": "2026-12-16", "Заметки": ""},
]

# What a tender usually needs; matched against "Документ"/"Тип" by keywords.
REQUIRED_FOR_TENDER = [
    ("Свидетельство о регистрации", ("свидетельств", "гувоҳнома")),
    ("Лицензия на перевозки", ("лицензи",)),
    ("Полисы ОСГО на машины", ("осго",)),
    ("Техпаспорта / список машин", ("техпаспорт", "список машин")),
    ("Справка об отсутствии налоговой задолженности", ("налог",)),
    ("Справка MIB (нет долгов по решениям суда)", ("mib", "мib", "принудительн")),
    ("ЭЦП директора", ("эцп", "эри")),
]

HELP = (
    "📁 <b>Документовед</b>\n\n"
    "<code>/documents</code> — сроки документов компании и чего не хватает для тендеров\n"
    "<code>/paket 518494</code> — черновик пакета на тендер (номер лота)\n"
    "<code>/paket 518494 21500000</code> — с ценой за рейс\n\n"
    "Реквизиты — лист «Реквизиты компании», документы — лист «Документы»."
)


# ---------------------------------------------------------------- pure helpers

def requisites_dict(rows: List[Dict[str, str]]) -> Dict[str, str]:
    out = {name: value for name, value, _ in DEFAULT_REQUISITES}
    for row in rows:
        name = (row.get("Параметр") or "").strip()
        if name:
            out[name] = (row.get("Значение") or "").strip()
    return out


def document_status(rows: List[Dict[str, str]], now: Optional[datetime] = None,
                    warn_days: int = 30) -> Dict[str, List[Tuple[str, str]]]:
    now = now or now_local()
    result: Dict[str, List[Tuple[str, str]]] = {"expired": [], "soon": [], "ok": [], "missing": []}
    for row in rows:
        name = row.get("Документ") or ""
        until = parse_date(row.get("Действует до"))
        if until is None:
            result["ok"].append((name, row.get("Заметки") or "без срока"))
            continue
        left = (until.date() - now.date()).days
        if left < 0:
            result["expired"].append((name, f"просрочен {until.strftime('%d.%m.%Y')}"))
        elif left <= warn_days:
            result["soon"].append((name, f"до {until.strftime('%d.%m.%Y')}, осталось {left} дн."))
        else:
            result["ok"].append((name, f"до {until.strftime('%d.%m.%Y')}"))
    text = " ".join(f"{r.get('Документ', '')} {r.get('Тип', '')}" for r in rows).lower()
    for title, keys in REQUIRED_FOR_TENDER:
        if not any(k in text for k in keys):
            result["missing"].append((title, "нет в листе «Документы»"))
    return result


def format_status(status: Dict[str, List[Tuple[str, str]]], requisites: Dict[str, str]) -> str:
    e = html.escape
    lines = ["📁 <b>Документы компании</b>", ""]
    for key, title in (("expired", "❌ Просрочены"), ("soon", "⚠️ Скоро истекают"),
                       ("missing", "📭 Нет для тендеров"), ("ok", "✅ В порядке")):
        if status[key]:
            lines.append(f"<b>{title}:</b>")
            lines += [f"   • {e(n)} — {e(d)}" for n, d in status[key]]
            lines.append("")
    empty = [k for k, v in requisites.items() if not v]
    if empty:
        lines.append("✏️ Не заполнено в «Реквизиты компании»: " + e(", ".join(empty)))
    return "\n".join(lines).strip()


def tender_from_row(row: Dict[str, str]) -> Dict[str, str]:
    def first(*names: str) -> str:
        for n in names:
            if row.get(n):
                return row[n].strip()
        return ""
    url = first("Ссылка")
    lot = first("Номер лота", "ID")
    if not lot:
        m = re.search(r"/lot/(\d+)", url)
        lot = m.group(1) if m else ""
    return {"lot": lot, "title": first("Название") or "[харид предмети]",
            "customer": first("Заказчик") or "[буюртмачи]", "amount": first("Сумма"),
            "deadline": first("Срок окончания", "Окончание"), "url": url}


def find_tender(tables: List[List[Dict[str, str]]], query: str) -> Optional[Dict[str, str]]:
    q = re.sub(r"\D", "", query or "")
    if not q:
        return None
    for rows in tables:
        for row in rows:
            fields = (row.get("Номер лота", ""), row.get("ID", ""), row.get("Ссылка", ""), row.get("Stable Key", ""))
            if any(q == re.sub(r"\D", "", f) or f"/lot/{q}" in f for f in fields if f):
                return tender_from_row(row)
    return None


def parse_paket(text: str) -> Tuple[str, Optional[float]]:
    body = re.sub(r"^/paket(@\w+)?", "", (text or "").strip(), flags=re.I).strip()
    m = re.search(r"/lot/(\d+)", body)
    tokens = body.replace("/lot/", " ").split()
    lot = m.group(1) if m else (re.sub(r"\D", "", tokens[0]) if tokens else "")
    price = None
    for tok in tokens[1:]:
        digits = re.sub(r"[^\d]", "", tok)
        if len(digits) >= 5:
            price = float(digits)
    return lot, price


def money(value: Optional[float]) -> str:
    return "[__________]" if not value else f"{value:,.0f}".replace(",", " ")


# ---------------------------------------------------------------- Word generation

def _docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm
    return Document, WD_ALIGN_PARAGRAPH, Pt, Cm


class Writer:
    """Tiny helper over python-docx: [text in brackets] is highlighted yellow."""

    def __init__(self):
        Document, self.ALIGN, self.Pt, Cm = _docx()
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = self.Pt(12)
        for section in self.doc.sections:
            section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
            section.top_margin = section.bottom_margin = Cm(1.8)

    def p(self, text: str = "", bold: bool = False, align: str = "justify", size: int = 12,
          indent: bool = False, after: int = 6) -> None:
        from docx.enum.text import WD_COLOR_INDEX
        par = self.doc.add_paragraph()
        par.alignment = {"justify": self.ALIGN.JUSTIFY, "center": self.ALIGN.CENTER,
                         "right": self.ALIGN.RIGHT, "left": self.ALIGN.LEFT}[align]
        par.paragraph_format.space_after = self.Pt(after)
        if indent:
            par.paragraph_format.first_line_indent = self.Pt(28)
        for part in re.split(r"(\[[^\]]+\])", text):
            if not part:
                continue
            run = par.add_run(part)
            run.bold = bold
            run.font.size = self.Pt(size)
            if part.startswith("["):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def table(self, rows: List[List[str]]) -> None:
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(value))
                run.font.size = self.Pt(10)
                run.bold = i == 0
        self.doc.add_paragraph()

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def _head(w: Writer, r: Dict[str, str]) -> None:
    w.p(f"Иштирокчи: {r['Название']}", after=0)
    w.p(f"СТИР: {r['СТИР'] or '[СТИР]'}   Манзил: {r['Адрес'] or '[манзил]'}", after=0)
    w.p(f"Банк: {r['Банк'] or '[банк]'}  МФО: {r['МФО'] or '[МФО]'}  Ҳ/р: {r['Расчётный счёт'] or '[ҳисоб рақами]'}", after=0)
    w.p(f"Тел.: {r['Телефон']}   E-mail: {r['E-mail']}", after=12)


def _sign(w: Writer, r: Dict[str, str]) -> None:
    w.p("")
    w.p(f"{r['Название']} директори", bold=True, after=0)
    w.p(f"______________________  {r['Директор']}", after=0)
    w.p("              (имзо)                                         М.Ў.", size=10)
    w.p(f"Сана: «[__]» {now_local().strftime('%m.%Y')} й.")


def _lot_text(t: Dict[str, str]) -> str:
    return f"«{t['title']}» (лот № {t['lot']})"


def build_package(tender: Dict[str, str], r: Dict[str, str], price: Optional[float] = None,
                  trips: Optional[int] = None) -> List[Tuple[str, bytes]]:
    lot = _lot_text(tender)
    cust = tender["customer"]
    files: List[Tuple[str, bytes]] = []

    # 1. cover letter
    w = Writer(); _head(w, r)
    w.p(cust, bold=True, align="right", after=0); w.p("Харид комиссиясига", align="right", after=12)
    w.p("ТАКЛИФ ХАТИ", bold=True, align="center", size=14)
    w.p(f"{r['Название']} (кейинги ўринларда — Иштирокчи) {cust} томонидан эълон қилинган {lot} харид ҳужжатлари, "
        "шу жумладан техник топшириқ билан тўлиқ танишиб чиқиб, қуйидаги таклифни тақдим этади.", indent=True)
    w.p("1. Таклиф нархи:", indent=True)
    w.table([["Хизмат номи", "Ҳажм", "Бирлик нархи, сўм (ҚҚС билан)", "Жами, сўм (ҚҚС билан)"],
             [tender["title"], f"{trips} рейс" if trips else "[ҳажм]", money(price),
              money(price * trips) if price and trips else "[__________]"]])
    w.p("Нархга транспорт воситасидан фойдаланиш, ҳайдовчи меҳнати, ёқилғи, юкни маҳкамлаш воситалари ва шартнома "
        "доирасида назарда тутилган бошқа барча харажатлар киритилган.", indent=True)
    w.p("2. Иштирокчи харид ҳужжатларида белгиланган тўлов шартлари ва хизмат кўрсатиш муддатларига розилик билдиради.", indent=True)
    w.p("3. Мазкур таклиф харид ҳужжатларида белгиланган муддат давомида амал қилади. Иштирокчи тақдим этилаётган "
        "маълумот ва ҳужжатларнинг ҳақиқийлиги учун жавобгарликни ўз зиммасига олади.", indent=True)
    _sign(w, r); files.append(("1_Таклиф_хати.docx", w.bytes()))

    # 2. annex 6 anti-corruption
    w = Writer()
    w.p("Харид қилиш ҳужжатларига 6-илова", align="right", size=10, after=12)
    w.p("КОРРУПЦИЯ КЎРИНИШЛАРИГА ЙЎЛ ҚЎЙМАСЛИК БЎЙИЧА АРИЗА", bold=True, align="center")
    w.p(f"Мен, {r['Директор']}, {r['Название']} директори, Устав асосида иш юритувчи, {cust} томонидан эълон "
        f"қилинган {lot} савдоларида иштирок этиш доирасида Иштирокчи номидан қуйидагиларни кафолатлайман:", indent=True)
    for i, item in enumerate([
        "Иштирокчи, унинг раҳбарлари ва ходимлари харид комиссияси аъзоларига, буюртмачининг мансабдор шахслари ва ходимларига бевосита ёки билвосита пора бермаган, ваъда қилмаган ва таклиф этмаган.",
        "Иштирокчи савдо натижаларига қонунга хилоф равишда таъсир кўрсатиш мақсадида ҳеч қандай моддий ёки номоддий наф, имтиёз ва устунликлар таклиф этмайди.",
        "Иштирокчи бошқа иштирокчилар билан нархларни келишиш ёки рақобатни чеклаш бўйича келишувларга кирмаган.",
        "Иштирокчи буюртмачи, харид комиссияси аъзолари ва ушбу лотдаги бошқа иштирокчилар билан аффилланган эмас, манфаатлар тўқнашуви мавжуд эмас.",
        "Коррупция ҳолатлари маълум бўлганда Иштирокчи буюртмачини ва ваколатли давлат органларини дарҳол хабардор қилади.",
        "Мажбуриятлар бузилган ва бу ваколатли орган томонидан тасдиқланган тақдирда таклиф рад этилиши ва шартнома компенсациясиз бекор қилинишига розилик билдираман.",
    ], 1):
        w.p(f"{i}. {item}", indent=True)
    _sign(w, r); files.append(("2_Илова_6_Коррупция.docx", w.bytes()))

    # 3. affiliation
    w = Writer(); _head(w, r)
    w.p("МАНФААТЛАР ТЎҚНАШУВИ ВА АФФИЛЛАНГАНЛИК ЙЎҚЛИГИ ТЎҒРИСИДА АРИЗА", bold=True, align="center")
    w.p(f"{r['Название']} {lot} савдоларида иштирок этиш доирасида маълум қилади: Иштирокчи, унинг муассислари ва "
        f"раҳбарлари {cust}, унинг мансабдор шахслари, харид комиссияси аъзолари ва ушбу лотдаги бошқа иштирокчилар "
        "билан аффилланган эмас; манфаатлар тўқнашуви мавжуд эмас; Иштирокчи банкротлик тартиб-таомилида эмас, солиқлар "
        "бўйича муддати ўтган қарздорлиги ва суд қарорлари бўйича бажарилмаган мажбуриятлари йўқ, инсофсиз "
        "ижрочиларнинг ягона реестрига киритилмаган.", indent=True)
    _sign(w, r); files.append(("3_Аффилланганлик.docx", w.bytes()))

    # 4. technical proposal (generic part; specifics of the ТЗ added by a person)
    w = Writer(); _head(w, r)
    w.p("ТЕХНИК ТАКЛИФ", bold=True, align="center", size=14)
    w.p(lot, align="center", after=12)
    w.p("1. Иштирокчи ҳақида маълумот", bold=True)
    w.p(f"Гувоҳнома {r['Свидетельство']}. Лицензия {r['Лицензия']}. {r['Название']} 2015 йилдан бери юкларни "
        "автомобиль транспортида ташиш хизматларини кўрсатиб келмоқда; шартномаларнинг асосий қисми очиқ танлов ва "
        "тендерлар натижасида тузилган.", indent=True)
    w.table([["Буюртмачи", "Ҳамкорлик даври"],
             ["«Ўзбектелеком» АК", "2018 йилдан ҳозиргача"], ["Uzmobile", "2016–2022 йй."],
             ["«Таъминот ва логистика хизмати» МЧЖ", "2023 йилдан ҳозиргача"],
             ["БТСЁР лойиҳаси ЛАГ (Осиё тараққиёт банки)", "2025 йилдан ҳозиргача"],
             ["UzMedImpeks", "2019–2023 йй."], ["Mobiuz", "2022–2024 йй."], ["«Навоийуран» ДК", "2024 й."]])
    w.p("2. Техник топшириқ талабларига мувофиқлик", bold=True)
    w.p(f"Ўз транспорт воситалари: {r['Парк']}. Паркда мавжуд бўлмаган транспорт турлари ҳамкор ташувчилар орқали "
        f"тақдим этилади: {r['Партнёры (самосвалы, фургоны, манипуляторы)'] or '[ҳамкорлар]'}.", indent=True)
    w.p(f"Рейс олди техник кўриги ва ҳайдовчиларнинг тиббий кўриги: {r['Техосмотр и медосмотр']}. Автотураргоҳ: "
        f"{r['Стоянка']}. Ҳайдовчилар C, CE тоифасига ва камида {r['Опыт водителей (лет)']} йиллик тажрибага эга.", indent=True)
    w.p("[Техник топшириқнинг алоҳида талаблари бўйича жавоблар — ҳажм, юк турлари, транспорт турлари]", indent=True)
    w.p("3. Суғурталаш", bold=True)
    w.p(f"Барча транспорт воситалари бўйича ОСГО ВТС полислари амал қилади; шартнома муддати давомида юклар "
        f"суғурталанади. Суғурта ташкилоти: {r['Страховая компания'] or '[суғурта ташкилоти]'}.", indent=True)
    w.p("4. Хизмат кўрсатиш муддатлари", bold=True)
    w.p(f"Шартнома тузилгандан сўнг {r['Начало работ после договора (рабочих дней)']} иш куни ичида хизмат кўрсатишга "
        f"киришилади. Буюртма бўйича транспорт Тошкент бўйича {r['Подача машины по Ташкенту (часов)']} соат, вилоятлар "
        f"бўйича {r['Подача машины по регионам (часов)']} соат ичида етказилади. Буюртмалар: {r['Телефон']}.", indent=True)
    _sign(w, r); files.append(("4_Техник_таклиф.docx", w.bytes()))

    # 5. annex 7 comparison (frame to fill from the ТЗ)
    w = Writer()
    w.p("Харид қилиш ҳужжатларига 7-илова", align="right", size=10, after=12)
    w.p("ТЕХНИК ТОПШИРИҚ ТАЛАБЛАРИ БИЛАН ТЕХНИК ЖИҲАТДАН ҚИЁСЛАШ ЖАДВАЛИ", bold=True, align="center")
    w.p(lot, align="center")
    rows = [["№", "Техник топшириқ талаби", "Иштирокчи таклифи", "Мувофиқлиги"]]
    generic = [
        ("Хизмат ҳажми (рейслар сони) буюртмалар асосида", "Ўз парки ва ҳамкорлар ҳисобидан тўлиқ таъминланади"),
        ("Юк турига мос транспорт воситаси", "Трал 1,5–50 т, тент 96 м³, контейнер ташувчи, бортли; бошқалари — ҳамкорлар"),
        ("Транспорт воситалари техник соз, техник кўрикдан ўтган", f"Ҳар куни рейс олди кўриги — {r['Техосмотр и медосмотр']}"),
        ("Ҳайдовчилар малакаси ва тиббий кўриги", f"C, CE тоифалари, {r['Опыт водителей (лет)']}+ йил тажриба, ҳар куни тиббий кўрик"),
        ("Юкни маҳкамлаш ва сақлаш", "Тасма, трос, занжир, ёпқичлар; юк суғурталанади"),
        ("Транспортни ўз вақтида етказиш", f"Тошкент — {r['Подача машины по Ташкенту (часов)']} соат, вилоятлар — {r['Подача машины по регионам (часов)']} соат"),
        ("Ҳужжатлар: ТТН, йўл варақаси, далолатнома", "Таъминланади"),
        ("Меҳнат муҳофазаси, ёнғин ва йўл хавфсизлиги", "Таъминланади; тиббий қути, ўт ўчиргич мавжуд"),
    ]
    for i, (req, offer) in enumerate(generic, 1):
        rows.append([str(i), req, offer, "Мувофиқ"])
    rows.append([str(len(generic) + 1), "[техник топшириқнинг бошқа бандлари]", "[ ]", "[ ]"])
    w.table(rows)
    _sign(w, r); files.append(("5_Илова_7_Қиёслаш_жадвали.docx", w.bytes()))
    return files


def checklist_text(tender: Dict[str, str], status: Dict[str, List[Tuple[str, str]]], price: Optional[float]) -> str:
    e = html.escape
    lines = [f"📦 <b>Черновик пакета: лот {e(tender['lot'])}</b>",
             f"{e(tender['title'])}", f"Заказчик: {e(tender['customer'])}"]
    if tender.get("amount"):
        lines.append(f"Стартовая сумма: {e(tender['amount'])}")
    if tender.get("deadline"):
        lines.append(f"⏰ Подача до: <b>{e(tender['deadline'])}</b>")
    if tender.get("url"):
        lines.append(e(tender["url"]))
    lines += ["", "<b>Что сделать:</b>",
              "1. Прочитать ТЗ и дописать в «Техник таклиф» и «7-илова» пункты ТЗ (жёлтые поля)",
              "2. Цена за рейс" + (f": {money(price)} сум — проверить" if price else ": /price … и вписать"),
              "3. Задаток на лицевой счёт ҲККП (обычно 3% от суммы) — проверить в объявлении",
              "4. Проверить отсутствие долгов по налогам и в MIB",
              "5. Подписать, печать, PDF, загрузить и подписать заявку ЭЦП"]
    problems = status["expired"] + status["soon"] + status["missing"]
    if problems:
        lines += ["", "<b>⚠️ Документы:</b>"] + [f"   • {e(n)} — {e(d)}" for n, d in problems]
    return "\n".join(lines)


# ---------------------------------------------------------------- agent

class DocumentAgent:
    def __init__(self, base_factory: Callable[[], Optional[Base]]):
        self.base_factory = base_factory

    def _base(self) -> Base:
        base = self.base_factory()
        if base is None:
            raise RuntimeError("Google Sheets не настроен")
        return base

    def setup(self, base: Optional[Base] = None) -> Dict[str, int]:
        """Seed requisites and known company documents; never overwrites existing rows."""
        base = base or self._base()
        added: Dict[str, int] = {}
        known = {(r.get("Параметр") or "").strip() for r in base.records(REQ_SHEET)}
        new = [{"Параметр": n, "Значение": v, "Описание": d} for n, v, d in DEFAULT_REQUISITES if n not in known]
        if new:
            added[REQ_SHEET] = base.append_many(REQ_SHEET, new)
        docs = {(r.get("Документ") or "").strip() for r in base.records(DOC_SHEET)}
        new_docs = [dict(d, **{"Ответственный": "Тендерный специалист"}) for d in DEFAULT_DOCUMENTS
                    if d["Документ"] not in docs]
        if new_docs:
            added[DOC_SHEET] = base.append_many(DOC_SHEET, new_docs)
        if added:
            base.log(AGENT, "Созданы реквизиты и документы", ", ".join(f"{k}: {v}" for k, v in added.items()))
        return added

    def status_text(self) -> str:
        base = self._base()
        return format_status(document_status(base.records(DOC_SHEET)),
                             requisites_dict(base.records(REQ_SHEET)))

    def _tender_tables(self, base: Base) -> List[List[Dict[str, str]]]:
        tables = []
        titles = set(base.existing_titles())
        for name in TENDER_SHEETS:
            if name in titles:
                try:
                    values = base.spreadsheet.worksheet(name).get_all_values()
                except Exception:
                    continue
                if values:
                    headers = [str(h).strip() for h in values[0]]
                    tables.append([{h: (row[i].strip() if i < len(row) else "") for i, h in enumerate(headers) if h}
                                   for row in values[1:]])
        return tables

    def package(self, text: str) -> Tuple[str, List[Tuple[str, bytes]]]:
        base = self._base()
        lot, price = parse_paket(text)
        if not lot:
            return HELP, []
        tender = find_tender(self._tender_tables(base), lot) or {
            "lot": lot, "title": "[харид предмети]", "customer": "[буюртмачи]", "amount": "",
            "deadline": "", "url": f"https://etender.uzex.uz/lot/{lot}"}
        requisites = requisites_dict(base.records(REQ_SHEET))
        files = build_package(tender, requisites, price)
        status = document_status(base.records(DOC_SHEET))
        deadline = parse_date(tender.get("deadline"))
        due = (deadline - timedelta(days=1)).strftime("%Y-%m-%d") if deadline else ""
        base.create_task(f"Подать заявку: лот {lot} — {tender['title'][:80]}", agent=AGENT,
                         assignee="Тендерный специалист", due=due, link=tender.get("url", ""),
                         notes="Черновик пакета отправлен в группу; проверить, подписать, загрузить.")
        base.log(AGENT, "Черновик пакета на тендер", f"лот {lot}, файлов {len(files)}", lot)
        return checklist_text(tender, status, price), files
