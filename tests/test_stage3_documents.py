"""Stage 3 tests: Document specialist on the in-memory spreadsheet."""

import io
from datetime import datetime

from trans_ocean.base import TZ, Base
from trans_ocean.documents import (DocumentAgent, build_package, document_status, find_tender,
                                   parse_paket, requisites_dict)
from tests.test_stage0 import FakeSpreadsheet, FakeWorksheet


def make(with_tender=True):
    sheets = []
    if with_tender:
        sheets.append(FakeWorksheet("Тендеры", [
            ["ID", "Номер лота", "Название", "Заказчик", "Сумма", "Срок окончания", "Ссылка"],
            ["518494", "518494", "TURLI XIL YUKLARNI AVTOMOBIL TRANSPORTIDA TASHISH", '"NEFTGAZMONTAJ" MCHJ',
             "1792040000", "29.10.2026 12:01", "https://etender.uzex.uz/lot/518494"],
        ]))
    base = Base(FakeSpreadsheet(sheets))
    base.ensure_sheets()
    agent = DocumentAgent(lambda: base)
    agent.setup(base)
    return base, agent


def docx_text(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_setup_seeds_once():
    base, agent = make()
    req = requisites_dict(base.records("Реквизиты компании"))
    assert req["СТИР"] == "303372643" and "Тўлаган" in req["Директор"]
    assert len(base.records("Документы")) == 6
    assert agent.setup(base) == {}


def test_document_status():
    base, _ = make()
    status = document_status(base.records("Документы"), now=datetime(2026, 11, 20, tzinfo=TZ))
    soon = dict(status["soon"])
    assert any("стоянки" in n for n in soon)                      # 16.12.2026 -> 26 days
    missing = [n for n, _ in status["missing"]]
    assert "Полисы ОСГО на машины" in missing and "Свидетельство о регистрации" not in missing
    expired = document_status(base.records("Документы"), now=datetime(2027, 1, 5, tzinfo=TZ))["expired"]
    assert len(expired) == 3


def test_parse_and_find():
    assert parse_paket("/paket 518494") == ("518494", None)
    assert parse_paket("/paket https://etender.uzex.uz/lot/518494 21 500 000".replace("21 500 000", "21500000")) == ("518494", 21500000.0)
    rows = [{"Ссылка": "https://etender.uzex.uz/lot/518494", "Заказчик": "X", "Название": "Y"}]
    t = find_tender([rows], "518494")
    assert t["lot"] == "518494" and t["customer"] == "X"
    assert find_tender([rows], "1") is None


def test_package_files_contain_requisites_and_price():
    base, agent = make()
    text, files = agent.package("/paket 518494 21500000")
    assert len(files) == 5 and "NEFTGAZMONTAJ" in text and "29.10.2026" in text
    cover = docx_text(dict(files)["1_Таклиф_хати.docx"])
    assert "303372643" in cover and "21 500 000" in cover and "NEFTGAZMONTAJ" in cover
    annex6 = docx_text(dict(files)["2_Илова_6_Коррупция.docx"])
    assert "Нуриддинов Тўлаган Акоидинович" in annex6
    tasks = [t for t in base.records("Задачи") if t["Агент"] == "Документовед"]
    assert len(tasks) == 1 and tasks[0]["Срок"] == "2026-10-28"


def test_package_without_tender_row_still_builds():
    base, agent = make(with_tender=False)
    text, files = agent.package("/paket 777777")
    assert len(files) == 5 and "777777" in text
    assert "Полисы ОСГО" in text                                    # missing docs listed
